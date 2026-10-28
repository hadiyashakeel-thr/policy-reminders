"""
Sends scheduled HR policy reminders to Discord channels via webhooks.

Run once a day, shortly after 7:00 PM Pakistan time (PKT = UTC+5, no DST).
It looks up today's policy for each group in Policy_Schedules_2.xlsx and posts it.

Usage:
    python policy_sender.py              # live run
    python policy_sender.py --dry-run    # print only, send nothing
    python policy_sender.py --dry-run --date 2026-11-04   # simulate a date
"""
import argparse
import json
import os
import re
import sys
import urllib.request
from datetime import date, datetime, timedelta, timezone

import openpyxl

# ----------------------------- CONFIG ---------------------------------
XLSX_PATH = "Policy_Schedules_2.xlsx"
PKT = timezone(timedelta(hours=5))
SEND_AFTER_HOUR = 19                 # 7:00 PM PKT
# Monthly cycle: Week 1 = days 1-7, Week 2 = days 8-14, ... Week 5 = days 29-31.
# The same schedule repeats every month.

# Discord group -> which column block of the sheet it follows.
# Blocks: "PRESTIGE" (cols A-C), "TESTFIESTA_LENS" (cols E-G), "QUALALIGN" (cols I-K)
# DEMAND CREW has no schedule in the sheet yet, so it is None (skipped).
GROUPS = {
    "PRESTIGE":    "PRESTIGE",
    "TEST FIESTA": "TESTFIESTA_LENS",
    "LENS":        "TESTFIESTA_LENS",
    "QUAL ALIGN":  "QUALALIGN",
    "DEMAND CREW": None,
}
# Webhook URLs come from environment variables (never hard-code them):
#   WEBHOOK_PRESTIGE, WEBHOOK_TEST_FIESTA, WEBHOOK_LENS,
#   WEBHOOK_QUAL_ALIGN, WEBHOOK_DEMAND_CREW
# ----------------------------------------------------------------------

BLOCK_COLS = {"PRESTIGE": 3, "TESTFIESTA_LENS": 7, "QUALALIGN": 11}  # policy-name column (1-based)
DAY_OFFSET = {"Monday": 0, "Tuesday": 1, "Wednesday": 2, "Thursday": 3, "Friday": 4}


def clean(text):
    return re.sub(r"\s+", " ", str(text).replace("\xa0", " ")).strip()


def load_schedule(path):
    """Returns {block: {(week, weekday_name): policy_name}}"""
    ws = openpyxl.load_workbook(path).active
    schedule = {b: {} for b in BLOCK_COLS}
    for block, pcol in BLOCK_COLS.items():
        week = None
        for row in range(1, ws.max_row + 1):
            label = ws.cell(row, pcol - 2).value  # Sr. No / "Week # N" column
            day = ws.cell(row, pcol - 1).value
            policy = ws.cell(row, pcol).value
            if isinstance(label, str):
                m = re.match(r"Week\s*#\s*(\d+)", label.strip(), re.I)
                if m:
                    week = int(m.group(1))
                    continue
            if week and day and policy:
                schedule[block][(week, clean(day))] = clean(policy)
    return schedule


def post(webhook, content):
    req = urllib.request.Request(
        webhook,
        data=json.dumps({"content": content}).encode(),
        headers={"Content-Type": "application/json", "User-Agent": "policy-sender/1.0"},
    )
    urllib.request.urlopen(req, timeout=15).read()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--date", help="simulate a date (YYYY-MM-DD); skips the 7 PM check")
    args = ap.parse_args()

    now = datetime.now(PKT)
    if args.date:
        today = date.fromisoformat(args.date)
    else:
        today = now.date()
        if now.hour < SEND_AFTER_HOUR and not args.dry_run:
            sys.exit(f"It is {now:%H:%M} PKT - policies go out after 7:00 PM PKT. Nothing sent.")

    weekday = today.strftime("%A")
    if weekday not in DAY_OFFSET:
        print(f"{weekday}: weekend, nothing scheduled.")
        return
    week = (today.day - 1) // 7 + 1
    schedule = load_schedule(XLSX_PATH)
    print(f"{today} ({weekday}) -> Week {week}")

    for group, block in GROUPS.items():
        if block is None:
            print(f"[{group}] no schedule defined - skipped")
            continue
        policy = schedule[block].get((week, weekday))
        if not policy:
            print(f"[{group}] no policy scheduled today")
            continue
        message = f"📢 **Policy Reminder**\n**{policy}**"
        env = "WEBHOOK_" + group.replace(" ", "_")
        if args.dry_run:
            print(f"[{group}] (dry run) {policy}")
            continue
        url = os.environ.get(env)
        if not url:
            print(f"[{group}] missing env var {env} - skipped")
            continue
        post(url, message)
        print(f"[{group}] sent: {policy}")


if __name__ == "__main__":
    main()
